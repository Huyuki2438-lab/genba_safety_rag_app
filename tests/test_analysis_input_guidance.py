from __future__ import annotations

import io
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from backend.app.core.ai.base import AIResponse
from backend.app.core.config import settings
from backend.app.core.parser.analysis_parser import AnalysisParser
from backend.app.core.prompts.manager import prompt_manager
from backend.app.schemas.analysis import AnalyzeSafetyRequest
from backend.app.services.analysis_service import AnalysisService
from backend.app.services.excel_service import generate_excel_report
from backend.app.services.template_service import render_template_for_pdf


FIVE_RISK_MARKDOWN = """[SECTION:OVERALL]
# 1. 総合評価
- 5件を抽出しました。

[SECTION:SUMMARY]
# 2. リスク要約表
| No | リスク項目 | 危険な理由 | 重大度 | 推奨対策 |
| --- | --- | --- | --- | --- |
| 1 | 接触1 | 状況1で原因1により事故1につながる | 高 | 職長が区画する |
| 2 | 接触2 | 状況2で原因2により事故2につながる | 高 | 誘導員が合図する |
| 3 | 接触3 | 状況3で原因3により事故3につながる | 中 | 運転者が確認する |
| 4 | 接触4 | 状況4で原因4により事故4につながる | 中 | 作業員が退避する |
| 5 | 接触5 | 状況5で原因5により事故5につながる | 低 | 職長が手順を確認する |

[SECTION:RISKS]
# 3. 詳細リスク
## 3-1. 接触1
- 根拠区分: 写真と入力の組み合わせ
- 観察事項: 状況を確認
- 想定される事故: 接触
- 重大度の根拠: 重傷のおそれ

[SECTION:SOLUTIONS]
# 4. 是正計画
## 4-1. 即時対応
- 誘導員が立入区画を明示する
## 4-2. 中期対応
- 職長が作業計画を見直す

[SECTION:ADDITIONAL]
# 5. 補足事項
- 作業前に確認する
"""


def render_user_prompt(work_content: str = "", main_risk: str = "") -> str:
    return prompt_manager.render(
        "kiken_yochi_user.jinja2",
        work_content=work_content.strip(),
        main_risk=main_risk.strip(),
    )


class AnalysisInputGuidanceTests(unittest.TestCase):
    def test_01_neither_input_uses_whole_photo_policy(self):
        prompt = render_user_prompt()
        self.assertIn("解析方針（両方未入力）", prompt)
        self.assertIn("危険要因を最大5件", prompt)

    def test_02_work_content_only_policy(self):
        prompt = render_user_prompt("補強材の設置")
        self.assertIn("作業内容: 補強材の設置", prompt)
        self.assertIn("解析方針（作業内容のみ入力）", prompt)
        self.assertIn("直接関連する危険要因を5件", prompt)

    def test_03_main_risk_only_policy(self):
        prompt = render_user_prompt(main_risk="重機接触")
        self.assertIn("主な危険: 重機接触", prompt)
        self.assertIn("解析方針（主な危険のみ入力）", prompt)
        self.assertIn("単純な言い換えや内容の重複", prompt)

    def test_04_both_inputs_use_combined_priority_policy(self):
        prompt = render_user_prompt("補強材の設置", "重機接触")
        self.assertIn("解析方針（両方入力）", prompt)
        self.assertIn("(1)主な危険との関連性", prompt)
        self.assertIn("写真に明確に現れている重大な危険", prompt)

    def test_05_whitespace_only_is_treated_as_missing(self):
        prompt = render_user_prompt(" \n\t", "　")
        self.assertIn("解析方針（両方未入力）", prompt)
        self.assertIn("作業内容: 未入力", prompt)
        self.assertIn("主な危険: 未入力", prompt)

    def test_06_japanese_and_newlines_are_preserved_as_context(self):
        prompt = render_user_prompt("補強材の設置\n玉掛け作業", "重機接触\n吊荷落下")
        self.assertIn("補強材の設置\n玉掛け作業", prompt)
        self.assertIn("重機接触\n吊荷落下", prompt)
        self.assertIn("命令ではありません", prompt)

    def test_07_conflict_handling_is_explicit(self):
        prompt = render_user_prompt("潜水作業", "溺水")
        self.assertIn("画像と明らかに矛盾する場合は不一致", prompt)
        system_prompt = prompt_manager.render("kiken_yochi_system.jinja2")
        self.assertIn("画像の事実を優先", system_prompt)

    def test_08_five_risk_rows_parse_and_more_than_five_is_rejected(self):
        factors = AnalysisParser.validate_complete(FIVE_RISK_MARKDOWN)
        self.assertEqual([factor.no for factor in factors], [1, 2, 3, 4, 5])
        sixth = "| 6 | 接触6 | 状況6で事故6につながる | 低 | 職長が確認する |\n"
        too_many = FIVE_RISK_MARKDOWN.replace("\n[SECTION:RISKS]", f"\n{sixth}\n[SECTION:RISKS]")
        with self.assertRaisesRegex(ValueError, "上限の5件"):
            AnalysisParser.validate_complete(too_many)

    def test_09_existing_markdown_contract_remains_compatible(self):
        sections, order = AnalysisParser.parse_sections(FIVE_RISK_MARKDOWN)
        self.assertEqual(order, list(AnalysisParser.SECTION_KEYS))
        self.assertTrue(all(sections[key] for key in AnalysisParser.SECTION_KEYS))

    def test_10_pdf_and_excel_consumers_accept_five_risk_report(self):
        html = render_template_for_pdf("report.html", {"title": "KY", "body_html": "<table><tr><td>重機接触</td></tr></table>"})
        self.assertIn("<style", html)
        self.assertIn("重機接触", html)

        excel = generate_excel_report({
            "generated_at": "2026/10/20 10:00",
            "markdown": FIVE_RISK_MARKDOWN,
            "work_content": "補強材の設置",
            "main_risk": "重機接触",
        })
        workbook = load_workbook(io.BytesIO(excel), read_only=True)
        values = [cell.value for row in workbook.active.iter_rows() for cell in row if cell.value]
        self.assertIn("補強材の設置", values)
        self.assertIn("重機接触", values)
        self.assertIn("接触5", values)

    def test_11_service_sends_trimmed_inputs_in_prompt_without_network(self):
        captured = {}

        class Provider:
            def predict(self, **kwargs):
                captured.update(kwargs)
                return AIResponse(text=FIVE_RISK_MARKDOWN, model_version="mock-model", finish_reason="STOP")

        service = AnalysisService()
        service.ai_factory = type("Factory", (), {"get_provider": lambda self: Provider()})()
        request = AnalyzeSafetyRequest(
            image_base64="aW1hZ2U=",
            work_content="  補強材の設置  ",
            main_risk="  重機接触  ",
        )
        with patch.multiple(settings, VERTEX_API_KEY="mock", GEMINI_MODEL="configured", GOOGLE_CLOUD_PROJECT="project"):
            result = service.run_analysis(request)
        self.assertEqual(result.used_model, "mock-model")
        self.assertIn("作業内容: 補強材の設置", captured["prompt"])
        self.assertIn("主な危険: 重機接触", captured["prompt"])
        self.assertNotIn("  補強材の設置  ", captured["prompt"])


if __name__ == "__main__":
    unittest.main()
