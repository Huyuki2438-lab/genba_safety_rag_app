from __future__ import annotations

import base64
import io
import logging
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import IllegalCharacterError
from PIL import Image as PillowImage

from backend.app.api.v1.excel import generate_excel
from backend.app.core.shared_storage import HistoryStorageUnavailable, publish
from backend.app.schemas.pdf import ExcelGenerateRequest
from backend.app.services.excel_service import _sanitize_excel_text, generate_excel_report
from desktop_app import _setup_logging


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)
RECORD_ID = "11111111-1111-4111-8111-111111111111"


def request(**overrides) -> ExcelGenerateRequest:
    values = {
        "generated_at": "2026/10/27 10:00",
        "markdown": "[SECTION:OVERALL]\n通常の日本語です。",
        "record_id": RECORD_ID,
        "site_name": "日本語の現場",
        "work_content": "掘削作業",
        "main_risk": "重機との接触",
        "created_by": "担当者",
    }
    values.update(overrides)
    return ExcelGenerateRequest(**values)


class ExcelTextSanitizationTests(unittest.TestCase):
    def test_removes_only_excel_forbidden_controls(self):
        value = "日本語\t記号！\n改行\r保持\x00削除\x08削除\x0b削除\x1f削除"
        self.assertEqual(
            _sanitize_excel_text(value),
            "日本語\t記号！\n改行\r保持削除削除削除削除",
        )

    def test_markdown_and_image_generate_a_readable_workbook(self):
        markdown = """[SECTION:OVERALL]
## **総合評価**
日本語\tタブ\n改行と `Markdown` \x00
[SECTION:SUMMARY]
| No | リスク項目 | 危険な理由 | 重大度 | 推奨対策 |
| --- | --- | --- | --- | --- |
| 1 | 転倒 | 足元不良 | 高 | 整理整頓 |
"""
        content = generate_excel_report(
            {
                "title": "現場安全\x01レポート",
                "generated_at": "2026/10/27 10:00",
                "markdown": markdown,
                "site_name": "日本語の現場\x02",
                "work_content": "掘削\t作業",
                "main_risk": "転倒\n接触",
                "created_by": "担当者",
                "image_data_url": "data:image/png;base64," + base64.b64encode(PNG_1X1).decode(),
            }
        )
        workbook = load_workbook(io.BytesIO(content))
        values = [cell.value for row in workbook.active.iter_rows() for cell in row if isinstance(cell.value, str)]
        joined = "\n".join(values)
        self.assertIn("日本語の現場", joined)
        self.assertIn("掘削\t作業", joined)
        self.assertIn("転倒\n接触", joined)
        self.assertNotIn("\x00", joined)
        self.assertNotIn("\x01", joined)
        self.assertNotIn("\x02", joined)
        self.assertEqual(len(workbook.active._images), 1)

    def test_mpo_image_is_normalized_before_workbook_save(self):
        first_frame = PillowImage.new("RGB", (4, 3), "red")
        second_frame = PillowImage.new("RGB", (4, 3), "blue")
        mpo = io.BytesIO()
        first_frame.save(mpo, format="MPO", save_all=True, append_images=[second_frame])

        content = generate_excel_report(
            {
                "markdown": "[SECTION:OVERALL]\nMPO画像の帳票です。",
                "image_data_url": "data:image/jpeg;base64," + base64.b64encode(mpo.getvalue()).decode(),
            }
        )

        workbook = load_workbook(io.BytesIO(content))
        self.assertEqual(len(workbook.active._images), 1)
        self.assertEqual(workbook.active._images[0].format, "png")

    def test_publish_supports_a_japanese_path(self):
        with tempfile.TemporaryDirectory(prefix="ky_excel_") as temp_dir:
            destination = Path(temp_dir) / "日本語 保存先" / "帳票.xlsx"
            publish(destination, b"PK test")
            self.assertEqual(destination.read_bytes(), b"PK test")


class ExcelApiErrorClassificationTests(unittest.TestCase):
    def test_history_read_error_is_sanitized_503(self):
        with patch(
            "backend.app.repositories.history_repository.history_repository.get",
            side_effect=HistoryStorageUnavailable(r"I:\機密\records"),
        ):
            with self.assertRaises(HTTPException) as raised:
                generate_excel(request())
        self.assertEqual(raised.exception.status_code, 503)
        self.assertIn("履歴を取得", raised.exception.detail)
        self.assertNotIn("I:", raised.exception.detail)

    def test_generation_error_is_sanitized_500(self):
        with (
            patch("backend.app.repositories.history_repository.history_repository.get", return_value=object()),
            patch(
                "backend.app.api.v1.excel.generate_excel_report",
                side_effect=IllegalCharacterError(r"secret C:\private"),
            ),
        ):
            with self.assertRaises(HTTPException) as raised:
                generate_excel(request())
        self.assertEqual(raised.exception.status_code, 500)
        self.assertIn("生成", raised.exception.detail)
        self.assertNotIn("secret", raised.exception.detail)

    def test_unwritable_destination_is_sanitized_503(self):
        with (
            patch("backend.app.repositories.history_repository.history_repository.get", return_value=object()),
            patch("backend.app.api.v1.excel.generate_excel_report", return_value=b"PK workbook"),
            patch("backend.app.core.shared_storage.require_root"),
            patch(
                "backend.app.core.shared_storage.publish",
                side_effect=HistoryStorageUnavailable(r"I:\機密\reports"),
            ),
        ):
            with self.assertRaises(HTTPException) as raised:
                generate_excel(request())
        self.assertEqual(raised.exception.status_code, 503)
        self.assertIn("保存", raised.exception.detail)
        self.assertNotIn("I:", raised.exception.detail)


class DesktopLoggingTests(unittest.TestCase):
    def test_api_exception_reaches_pc_local_log_handler(self):
        with tempfile.TemporaryDirectory(prefix="ky_logs_") as temp_dir:
            desktop_logger = _setup_logging(Path(temp_dir))
            app_logger = logging.getLogger("genba_safety_rag_app")
            try:
                try:
                    raise RuntimeError("diagnostic marker")
                except RuntimeError:
                    logging.getLogger("genba_safety_rag_app.api.excel").exception("Excel生成テスト")
                for handler in set(desktop_logger.handlers + app_logger.handlers):
                    handler.flush()
                log_text = next(Path(temp_dir).glob("*.log")).read_text(encoding="utf-8")
                self.assertIn("Excel生成テスト", log_text)
                self.assertIn("Traceback", log_text)
                self.assertIn("RuntimeError: diagnostic marker", log_text)
            finally:
                for logger in (desktop_logger, app_logger):
                    for handler in list(logger.handlers):
                        logger.removeHandler(handler)
                        handler.close()


if __name__ == "__main__":
    unittest.main()
