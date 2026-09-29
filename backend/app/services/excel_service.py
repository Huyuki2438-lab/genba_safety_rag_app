from __future__ import annotations

import base64
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Iterable


SECTION_TITLES = {
    "OVERALL": "1. 総合評価",
    "SUMMARY": "2. リスク要約表",
    "RISKS": "3. 詳細リスク",
    "SOLUTIONS": "4. 安全対策",
    "ADDITIONAL": "5. 補足事項",
}
SECTION_MARKER = re.compile(r"^\s*\[SECTION:([^\]]+)\]\s*$", re.IGNORECASE)
MARKDOWN_PREFIX = re.compile(r"^\s{0,3}(?:#{1,6}\s*|[-*+]\s+|\d+[.)]\s+|>\s?)")
MARKDOWN_EMPHASIS = re.compile(r"(\*\*|__|`|~~)")
SUBSECTION_PREFIX = re.compile(r"^\s*#{1,6}\s*(\d+[-－]\d+\.?\s*.*)$")
LABEL_VALUE = re.compile(r"^\s*[-*+]?\s*([^:：]+?)\s*[:：]\s*(.*)$")


@dataclass
class RiskDetail:
    title: str
    fields: list[tuple[str, str]] = field(default_factory=list)


@dataclass
class Measure:
    title: str
    rows: list[dict[str, str]] = field(default_factory=list)


def _plain_text(value: str) -> str:
    """Remove Markdown controls without changing the report's actual line breaks."""
    text = MARKDOWN_PREFIX.sub("", value).strip()
    text = MARKDOWN_EMPHASIS.sub("", text)
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)


def _section_lines(markdown: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {key: [] for key in SECTION_TITLES}
    active = "OVERALL"
    saw_marker = False
    for raw_line in markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        marker = SECTION_MARKER.match(raw_line)
        if marker:
            key = marker.group(1).upper()
            if key in sections:
                active = key
                saw_marker = True
            continue
        sections[active].append(raw_line)
    if not saw_marker:
        sections["OVERALL"] = markdown.replace("\r\n", "\n").split("\n")
    return sections


def _markdown_fields(value: str) -> list[str] | None:
    text = value.strip()
    if not (text.startswith("|") and text.endswith("|")):
        return None
    return [_plain_text(part) for part in text[1:-1].split("|")]


def _is_rule_row(fields: list[str]) -> bool:
    return bool(fields) and all(re.fullmatch(r":?-{3,}:?", field.strip()) for field in fields)


def _summary_rows(lines: list[str]) -> list[list[str]]:
    """Return the five columns in the risk summary table, one risk per row."""
    rows: list[list[str]] = []
    in_table = False
    for line in lines:
        fields = _markdown_fields(line)
        if not fields:
            if in_table:
                break
            continue
        if any("リスク項目" in field for field in fields):
            in_table = True
            continue
        if not in_table or _is_rule_row(fields) or len(fields) < 5:
            continue
        # Keep any unescaped extra pipe with the final recommendation instead
        # of shifting values into the wrong columns.
        row = fields[:4] + [" | ".join(fields[4:])]
        rows.append([value or "－" for value in row])
    return rows


def _non_table_text(lines: list[str]) -> list[str]:
    values: list[str] = []
    for line in lines:
        if _markdown_fields(line) is not None or line.lstrip().startswith("#"):
            continue
        if cleaned := _plain_text(line):
            values.append(cleaned)
    return values


def _risk_details(lines: list[str]) -> list[RiskDetail]:
    details: list[RiskDetail] = []
    current: RiskDetail | None = None
    current_field: int | None = None
    for raw_line in lines:
        if match := SUBSECTION_PREFIX.match(raw_line):
            current = RiskDetail(title=_plain_text(match.group(1)))
            details.append(current)
            current_field = None
            continue
        if raw_line.lstrip().startswith("#"):
            continue
        cleaned = _plain_text(raw_line)
        if not cleaned:
            continue
        if current is None:
            current = RiskDetail(title="詳細リスク")
            details.append(current)
        label_match = LABEL_VALUE.match(raw_line)
        if label_match:
            label = _plain_text(label_match.group(1))
            value = _plain_text(label_match.group(2)) or "－"
            current.fields.append((label, value))
            current_field = len(current.fields) - 1
        elif current_field is not None:
            label, value = current.fields[current_field]
            current.fields[current_field] = (label, f"{value}\n{cleaned}")
        else:
            current.fields.append(("内容", cleaned))
            current_field = len(current.fields) - 1
    return details


def _measure_values(value: str) -> dict[str, str]:
    """Extract a single countermeasure's four fields when labels are supplied."""
    normalized = _plain_text(value)
    aliases = {
        "対策": "対策", "対策名": "対策", "誰が": "誰が", "担当": "誰が",
        "何を": "何を", "実施内容": "何を", "どのように": "どのように", "方法": "どのように",
    }
    matches = list(re.finditer(r"(?:^|\s)(対策名?|誰が|担当|何を|実施内容|どのように|方法)\s*[:：]", normalized))
    values = {key: "－" for key in ("対策", "誰が", "何を", "どのように")}
    if not matches:
        values["対策"] = normalized or "－"
        return values
    for index, match in enumerate(matches):
        key = aliases[match.group(1)]
        end = matches[index + 1].start() if index + 1 < len(matches) else len(normalized)
        values[key] = normalized[match.end() : end].strip() or "－"
    return values


def _measures(lines: list[str]) -> list[Measure]:
    measures: list[Measure] = []
    current: Measure | None = None
    pending_row: dict[str, str] | None = None
    aliases = {
        "対策": "対策", "対策名": "対策", "誰が": "誰が", "担当": "誰が",
        "何を": "何を", "実施内容": "何を", "どのように": "どのように", "方法": "どのように",
    }
    for raw_line in lines:
        if match := SUBSECTION_PREFIX.match(raw_line):
            current = Measure(title=_plain_text(match.group(1)))
            measures.append(current)
            pending_row = None
            continue
        if raw_line.lstrip().startswith("#"):
            continue
        cleaned = _plain_text(raw_line)
        if not cleaned:
            continue
        if current is None:
            current = Measure(title="安全対策")
            measures.append(current)
        label_match = LABEL_VALUE.match(raw_line)
        label = _plain_text(label_match.group(1)) if label_match else ""
        label_value = _plain_text(label_match.group(2)) if label_match else ""
        if label in aliases:
            if pending_row is None or (label == "対策" and pending_row["対策"] != "－"):
                pending_row = {key: "－" for key in ("対策", "誰が", "何を", "どのように")}
                current.rows.append(pending_row)
            pending_row[aliases[label]] = label_value or "－"
            continue
        pending_row = _measure_values(cleaned)
        current.rows.append(pending_row)
    return measures


def _image_from_data_url(data_url: str):
    from openpyxl.drawing.image import Image
    header, separator, encoded = data_url.partition(",")
    if not separator or not header.lower().startswith("data:image/"):
        return None
    try:
        image = Image(io.BytesIO(base64.b64decode(encoded, validate=True)))
    except Exception:
        return None
    scale = min(1, 600 / image.width, 360 / image.height)
    image.width, image.height = int(image.width * scale), int(image.height * scale)
    return image


def _display_width(value: str) -> int:
    return sum(2 if ord(character) > 0xFF else 1 for character in value)


def _estimated_lines(value: object, characters_per_line: int) -> int:
    """Size rows without inserting artificial newlines into report content."""
    text = str(value or "－")
    return sum(max(1, (_display_width(line) + characters_per_line - 1) // characters_per_line) for line in text.splitlines() or [""])


def generate_excel_report(context: dict[str, object]) -> bytes:
    """Create an A4-printable workbook with real Excel tables for AI output."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "安全分析レポート"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = sheet.page_margins.right = 0.35
    sheet.page_margins.top = sheet.page_margins.bottom = 0.4
    sheet.freeze_panes = "A5"
    for column in range(1, 9):
        sheet.column_dimensions[get_column_letter(column)].width = 15

    navy, blue, pale_blue, white, dark, gray = "17365D", "D9EAF7", "EAF3F8", "FFFFFF", "1F2937", "F3F4F6"
    thin = Side(style="thin", color="AAB7C4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrapped = Alignment(horizontal="left", vertical="top", wrap_text=True)
    base_font = Font(name="Meiryo", size=10, color=dark)

    def merged_cell(row: int, first: int, last: int, value: object, *, fill: str | None = None,
                    font: Font | None = None, alignment: Alignment | None = None, height: float | None = None) -> None:
        sheet.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
        anchor = sheet.cell(row, first, value)
        anchor.font, anchor.alignment = font or base_font, alignment or wrapped
        if fill:
            anchor.fill = PatternFill("solid", fgColor=fill)
        for column in range(first, last + 1):
            cell = sheet.cell(row, column)
            cell.border = border
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
        if height is not None:
            sheet.row_dimensions[row].height = height

    def content_row(row: int, first: int, last: int, value: object, *, fill: str = pale_blue, chars: int = 75) -> None:
        height = max(22, min(409, _estimated_lines(value, chars) * 15 + 8))
        merged_cell(row, first, last, str(value or "－"), fill=fill, height=height)

    def section_header(row: int, value: str) -> None:
        merged_cell(row, 1, 8, value, fill=navy, font=Font(name="Meiryo", size=12, bold=True, color=white), alignment=Alignment(vertical="center"), height=24)

    def table_header(row: int, values: list[tuple[int, int, str]]) -> None:
        for first, last, value in values:
            merged_cell(row, first, last, value, fill=blue, font=Font(name="Meiryo", size=10, bold=True, color=dark), alignment=centered, height=28)

    title = str(context.get("title") or "現場安全 危険分析レポート")
    sheet.merge_cells("A1:H2")
    title_cell = sheet["A1"]
    title_cell.value, title_cell.font, title_cell.fill, title_cell.alignment = title, Font(name="Meiryo", size=18, bold=True, color=white), PatternFill("solid", fgColor=navy), centered
    for row in (1, 2):
        sheet.row_dimensions[row].height = 25
        for column in range(1, 9):
            sheet.cell(row, column).border = border
            sheet.cell(row, column).fill = PatternFill("solid", fgColor=navy)
    merged_cell(3, 1, 8, f"出力日時: {context.get('generated_at') or datetime.now().strftime('%Y/%m/%d %H:%M')}", fill=gray, alignment=Alignment(horizontal="right", vertical="center"), height=20)

    details = (("現場名", context.get("site_name")), ("作業内容", context.get("work_content")), ("主な危険", context.get("main_risk")), ("登録者", context.get("created_by")))
    for row, (label, value) in enumerate(details, start=5):
        text = str(value or "－")
        height = max(24, min(409, _estimated_lines(text, 75) * 15 + 8))
        merged_cell(row, 1, 2, label, fill=blue, font=Font(name="Meiryo", size=10, bold=True, color=dark), alignment=centered, height=height)
        merged_cell(row, 3, 8, text, fill=white, alignment=wrapped, height=height)

    next_row = 10  # exactly one spacer row after the metadata area
    image = _image_from_data_url(str(context.get("image_data_url") or ""))
    if image:
        section_header(next_row, "分析対象画像")
        next_row += 1
        sheet.add_image(image, f"A{next_row}")
        image_rows = max(14, int(image.height / 19) + 2)
        for row in range(next_row, next_row + image_rows):
            sheet.row_dimensions[row].height = 19
        next_row += image_rows + 1

    sections = _section_lines(str(context.get("markdown") or ""))

    section_header(next_row, SECTION_TITLES["OVERALL"])
    next_row += 1
    content_row(next_row, 1, 8, "\n".join(_non_table_text(sections["OVERALL"])) or "情報はありません。", chars=90)
    next_row += 2

    section_header(next_row, SECTION_TITLES["SUMMARY"])
    next_row += 1
    table_header(next_row, [(1, 1, "No"), (2, 3, "リスク項目"), (4, 5, "危険な理由"), (6, 6, "重大度"), (7, 8, "推奨対策")])
    next_row += 1
    summary_rows = _summary_rows(sections["SUMMARY"]) or [["－", "情報はありません。", "－", "－", "－"]]
    for values in summary_rows:
        height = max(26, min(409, max(_estimated_lines(values[0], 8), _estimated_lines(values[1], 28), _estimated_lines(values[2], 28), _estimated_lines(values[3], 8), _estimated_lines(values[4], 28)) * 15 + 8))
        for first, last, value, align in ((1, 1, values[0], centered), (2, 3, values[1], wrapped), (4, 5, values[2], wrapped), (6, 6, values[3], centered), (7, 8, values[4], wrapped)):
            merged_cell(next_row, first, last, value, fill=white, alignment=align, height=height)
        next_row += 1
    next_row += 1

    section_header(next_row, SECTION_TITLES["RISKS"])
    next_row += 1
    risk_details = _risk_details(sections["RISKS"]) or [RiskDetail("詳細リスク", [("内容", "情報はありません。")])]
    for detail in risk_details:
        merged_cell(next_row, 1, 8, detail.title, fill=blue, font=Font(name="Meiryo", size=11, bold=True, color=dark), height=22)
        next_row += 1
        table_header(next_row, [(1, 2, "項目"), (3, 8, "内容")])
        next_row += 1
        for label, value in detail.fields or [("内容", "情報はありません。")]:
            height = max(24, min(409, max(_estimated_lines(label, 18), _estimated_lines(value, 75)) * 15 + 8))
            merged_cell(next_row, 1, 2, label, fill=white, alignment=centered, height=height)
            merged_cell(next_row, 3, 8, value, fill=white, alignment=wrapped, height=height)
            next_row += 1
        next_row += 1

    section_header(next_row, SECTION_TITLES["SOLUTIONS"])
    next_row += 1
    measure_groups = _measures(sections["SOLUTIONS"]) or [Measure("安全対策", [{"対策": "情報はありません。", "誰が": "－", "何を": "－", "どのように": "－"}])]
    headers = [(1, 2, "対策"), (3, 3, "誰が"), (4, 5, "何を"), (6, 8, "どのように")]
    for group in measure_groups:
        merged_cell(next_row, 1, 8, group.title, fill=blue, font=Font(name="Meiryo", size=11, bold=True, color=dark), height=22)
        next_row += 1
        table_header(next_row, headers)
        next_row += 1
        for measure in group.rows or [{"対策": "情報はありません。", "誰が": "－", "何を": "－", "どのように": "－"}]:
            height = max(26, min(409, max(_estimated_lines(measure["対策"], 28), _estimated_lines(measure["誰が"], 12), _estimated_lines(measure["何を"], 28), _estimated_lines(measure["どのように"], 42)) * 15 + 8))
            for first, last, key in ((1, 2, "対策"), (3, 3, "誰が"), (4, 5, "何を"), (6, 8, "どのように")):
                merged_cell(next_row, first, last, measure[key], fill=white, alignment=wrapped, height=height)
            next_row += 1
        next_row += 1

    section_header(next_row, SECTION_TITLES["ADDITIONAL"])
    next_row += 1
    content_row(next_row, 1, 8, "\n".join(_non_table_text(sections["ADDITIONAL"])) or "情報はありません。", chars=90)

    sheet.print_title_rows = "1:3"
    sheet.print_area = f"A1:H{next_row}"
    sheet.oddFooter.center.text = "安全分析レポート  Page &P / &N"
    sheet.oddFooter.center.size = 8
    workbook.properties.creator = "現場安全分析アプリ"
    workbook.properties.title = title
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
